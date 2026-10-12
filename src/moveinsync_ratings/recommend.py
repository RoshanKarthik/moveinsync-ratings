"""The recommendation: what to fix and what to do first, built from `analyse()` and backlog.yaml.

backlog.yaml holds the judgement (the change, owner, effort, acceptance criteria). Every number in
the outputs comes from the analysis at render time, so nothing is typed by hand and a re-run
updates the evidence while the plan stays the same.

Outputs (in deliverables/):
- recommendation.html  answer-first decision document (one self-contained file)
- backlog.xlsx         the tickets with evidence, plus the issue evidence table
- backlog_jira.csv     Jira CSV import (Summary, Issue Type, Priority, Labels, Description)
- reply_playbook.md    reply templates and routing per issue, with the current reply backlog
"""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any, Literal

import pandas as pd
import yaml
from pydantic import BaseModel, Field

from . import charts
from .analysis import quotes_matching, two_prop_p, unanswered_negatives
from .charts import Row
from .evaluate import wilson
from .report import ROOT_VARS, _env, build_charts
from .taxonomy import Taxonomy

Horizon = Literal["now", "next", "quarter"]

HORIZONS: dict[str, tuple[str, str]] = {
    "now": ("This week", "No release needed"),
    "next": ("Next 1–2 releases", "Code changes, shipped in the next one or two app updates"),
    "quarter": ("This quarter", "Larger changes that need a sprint or more"),
}


# ── backlog model ─────────────────────────────────────────────────────────────────────────────
class Ticket(BaseModel):
    id: str
    title: str
    horizon: Horizon
    track: Literal["product", "ops"]
    type: str = "Task"
    priority: Literal["Highest", "High", "Medium", "Low"] = "Medium"
    owner: str
    effort: str
    issues: list[str] = []
    tracked: bool = True  # follow the first issue in the monitor's Fix tracker
    evidence: Literal["issue", "replies", "release", "releases", "silent", "safety", "orgs"] = "issue"
    quote_terms: list[str] = []
    depends_on: list[str] = []
    problem: str
    change: list[str]
    acceptance: list[str]
    metric: str
    confirm_with: str
    expected: str
    baseline: dict[str, Any] | None = None


class ReplyRule(BaseModel):
    issues: list[str]
    route: str
    template: str
    after_fix: str | None = None


class Backlog(BaseModel):
    version: int = 1
    tickets: list[Ticket]
    replies: list[ReplyRule] = Field(default_factory=list)


def load_backlog_model(path: Path, tax: Taxonomy) -> Backlog:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    backlog = Backlog.model_validate(data)
    ids = [t.id for t in backlog.tickets]
    if len(ids) != len(set(ids)):
        raise ValueError("backlog.yaml: ticket ids must be unique")
    for t in backlog.tickets:
        unknown = [c for c in t.issues if c not in tax.issues]
        if unknown:
            raise ValueError(f"backlog.yaml {t.id}: unknown issue code(s) {unknown}")
        missing = [d for d in t.depends_on if d not in ids]
        if missing:
            raise ValueError(f"backlog.yaml {t.id}: depends on unknown ticket(s) {missing}")
    for r in backlog.replies:
        unknown = [c for c in r.issues if c not in tax.issues]
        if unknown:
            raise ValueError(f"backlog.yaml replies: unknown issue code(s) {unknown}")
    return backlog


# ── baselines ─────────────────────────────────────────────────────────────────────────────────
def freeze_baselines(path: Path, a: dict[str, Any], rebaseline: bool = False) -> list[str]:
    """Write `baseline:` (recent share of 1–2★ reviews naming the ticket's first issue) into
    backlog.yaml for tickets that lack one. Edits the text in place so comments survive.
    Returns the ids that were (re)baselined."""
    by_code = {i["code"]: i for i in a["issues"]["issues"]}
    days = a["windows"]["recent"]["days"]
    as_of = a["end_iso"][:10]
    lines = path.read_text(encoding="utf-8").splitlines()
    data = yaml.safe_load("\n".join(lines))
    todo: dict[str, str] = {}
    for t in data.get("tickets", []):
        codes = [c for c in t.get("issues", []) if c in by_code]
        if not codes or t.get("tracked") is False or (t.get("baseline") and not rebaseline):
            continue
        s = by_code[codes[0]]["recent"]
        todo[t["id"]] = f"  baseline: {{share: {s['p']:.4f}, k: {s['k']}, n: {s['n']}, days: {days}, as_of: '{as_of}'}}"
    if not todo:
        return []
    out: list[str] = []
    current: str | None = None
    for line in lines:
        m = re.match(r"^- id:\s*(\S+)", line)
        if m:
            current = m.group(1)
        if current in todo and line.startswith("  baseline:"):
            continue  # replaced below
        out.append(line)
        if current in todo and line.startswith("  issues:"):
            out.append(todo[current])
    path.write_text("\n".join(out) + "\n", encoding="utf-8")
    return list(todo)


# ── evidence ──────────────────────────────────────────────────────────────────────────────────
def _p(v: float) -> str:
    """Whole percent, or one decimal under 10% so small shares and their intervals stay distinct."""
    return f"{v:.1%}" if 0 < v < 0.1 else f"{v:.0%}"


def _ci(s: dict[str, Any]) -> str:
    return f"{s['k']} of {s['n']}; 95% CI {_p(s['lo'])}–{_p(s['hi'])}"


def _pv(p: float) -> str:
    return f"p = {p:.3f}" if p >= 0.001 else "p < 0.001"


def detectable_share(baseline: float, n: int) -> float | None:
    """Largest share the monitor could call 'improved' at this sample size: the whole 95% Wilson
    interval must sit below the baseline. None when only zero mentions would do (too rare to confirm)."""
    best = None
    for k in range(1, n + 1):
        if wilson(k, n)[1] < baseline:
            best = k / n
        else:
            break
    return best


def issue_lines(i: dict[str, Any], a: dict[str, Any], emerging: dict[str, Any]) -> list[str]:
    days = a["windows"]["recent"]["days"]
    trend = (
        f"{i['trend']}, {_pv(i['trend_p'])}"
        if i["trend"] != "steady"
        else f"no significant change, {_pv(i['trend_p'])}"
    )
    lines = [
        f"{i['label']}: in {_p(i['recent']['p'])} of 1–2★ reviews in the last {days} days ({_ci(i['recent'])}); "
        f"{_p(i['prior']['p'])} in the {days} days before ({trend})."
    ]
    e = emerging.get(i["code"])
    if e:
        lines.append(
            f"Emerging: {_p(e['last30']['p'])} of 1–2★ reviews in the last 30 days ({_ci(e['last30'])}) vs "
            f"{_p(e['base']['p'])} in the 180 days before ({_pv(e['p'])})."
        )
    lines.append(
        f"Rating cost rank #{i['rank']} of {len(a['issues']['issues'])}: −{i['rating_cost']:.3f}★ on the "
        f"{a['windows']['year']['days']}-day written-review average; named in {i['mentions_year']} reviews in that period."
    )
    if i["blames_update"]["k"] >= 3 and i["blames_update"]["p"] >= 0.15:
        lines.append(f"{_p(i['blames_update']['p'])} of reviews naming it blame a recent update.")
    return lines


def ticket_evidence(t: Ticket, a: dict[str, Any], tax: Taxonomy) -> list[str]:
    by_code = {i["code"]: i for i in a["issues"]["issues"]}
    emerging = {e["code"]: e for e in a["emerging"]}
    W = a["windows"]
    days = W["recent"]["days"]
    out: list[str] = []
    if t.evidence == "replies":
        r = reply_story(a)
        out.append(
            f"Reviews answered: {_p(r['high']['replied']['p'])} in {r['high']['quarter']} (last quarter at 90% or more); "
            f"at most {_p(r['post_max'])} in every quarter since {r['collapse']}; last reply {a['replies']['last_reply']}."
        )
        out.append(
            f"Last {days} days: {_p(W['recent']['replied']['p'])} of reviews answered ({_ci(W['recent']['replied'])}); "
            f"1–2★ reviews answered: {_p(W['recent']['replied_neg']['p'])} ({_ci(W['recent']['replied_neg'])})."
        )
        for b in a["benchmark_reviews"][:1]:
            out.append(
                f"{b['name']}, last {W['year']['days']} days: answers {_p(b['replied']['p'])} of reviews and "
                f"{_p(b['replied_neg']['p'])} of 1–2★ reviews ({_ci(b['replied_neg'])})."
            )
        e = a["replies"]["edited_after_reply"]
        out.append(
            f"{e['k']:,} of {e['n']:,} reviews that got a reply ({_p(e['p'])}) were edited after it. Play keeps only "
            "the latest stars, so whether replies raise ratings is measured from now on (star-change log)."
        )
    elif t.evidence == "release":
        for v in [v for v in a["versions"] if v["flag"] in ("worse", "watch")][-2:]:
            top = ", ".join(f"{tax.issues[c].label} ({n})" for c, n in v["top_issues"])
            out.append(
                f"Release {v['version']} (first review {v['first_seen']}): {_p(v['neg']['p'])} of its {v['n']} reviews are 1–2★ "
                f"vs {_p(v['baseline_neg'])} for other releases ({_pv(v['p'])}; "
                + (
                    "significant after correction"
                    if v["flag"] == "worse"
                    else f"'watch' — not significant after correcting for {v['tests']} releases"
                )
                + f"). Most named in its 1–2★ reviews: {top}."
            )
    elif t.evidence == "releases":
        vs = a["versions"]
        flagged = [v for v in vs if v["flag"] in ("worse", "watch")]
        out.append(
            f"{len(vs)} releases had enough reviews to compare in the last {W['year']['days']} days; "
            f"{len(flagged)} flagged 'worse' or 'watch' by the monitor"
            + (f" (latest: {flagged[-1]['version']})." if flagged else ".")
        )
    elif t.evidence == "silent":
        g, m = a["gap"], a["listing"]["main"]
        written_1 = g["written"][0] / max(m["hist"][0], 1)
        out.append(
            f"{_p(1 - g['written_share_of_ratings'])} of the {m['ratings']:,} ratings have no text; they average "
            f"{g['star_only_mean']:.2f}★, vs {g['written_mean']:.2f}★ for written reviews (approximate: Google rescales its star counts)."
        )
        out.append(
            f"About {_p(written_1)} of the store's 1★ ratings come with a written review: low ratings are almost all "
            "written, while most ratings without text are 5★."
        )
    elif t.evidence == "safety":
        s = a["safety"]
        out.append(
            f"{s['n']} of {s['of']:,} reviews in the last {W['year']['days']} days raise a safety concern; "
            f"{s['last30']} in the last 30 days. {s['praise']} reviews praise safety."
        )
    elif t.evidence == "orgs":
        for o in a["orgs"][:5]:
            top = ", ".join(tax.issues[c].label for c, _ in o["top_issues"]) or "no specific issue"
            out.append(f"{o['org']}: named in {o['n']} reviews since 2023, {_p(o['neg']['p'])} 1–2★ ({top}).")
        if a["orgs"]:
            out.append("Small numbers: a lead for the account team, not a verdict.")
    # A release ticket cites its release plus its main issue; other tickets cite every issue.
    codes = t.issues[:1] if t.evidence == "release" else t.issues
    for code in codes:
        if code in by_code and t.evidence != "safety":
            out.extend(issue_lines(by_code[code], a, emerging))
    return out


def reply_story(a: dict[str, Any]) -> dict[str, Any]:
    """When replies stopped: last quarter at ≥90%, first quarter after it under 20%, and the
    highest rate in any quarter since."""
    q = [x for x in a["replies"]["quarters"]]
    high_idx = max((i for i, x in enumerate(q) if x["replied"]["p"] >= 0.9), default=0)
    after = q[high_idx + 1 :]
    collapse = next((x for x in after if x["replied"]["p"] < 0.2), after[0] if after else q[-1])
    since = q[q.index(collapse) :]
    return {
        "high": q[high_idx],
        "collapse": collapse["quarter"],
        "post_max": max(x["replied"]["p"] for x in since),
    }


# ── scenarios ─────────────────────────────────────────────────────────────────────────────────
def store_rating_if(neg_share: float, hist: list[int]) -> float:
    """Store rating if the 1–2★ share of ratings were `neg_share`, keeping today's mix within
    1–2★ and within 3–5★. Directional: Google's counts are rescaled to its weighted rating."""
    neg_n = hist[0] + hist[1]
    pos_n = sum(hist[2:])
    neg_mean = (hist[0] + 2 * hist[1]) / max(neg_n, 1)
    pos_mean = (3 * hist[2] + 4 * hist[3] + 5 * hist[4]) / max(pos_n, 1)
    return neg_share * neg_mean + (1 - neg_share) * pos_mean


def scenarios(a: dict[str, Any], backlog: Backlog) -> dict[str, Any]:
    m = a["listing"]["main"]
    hist = m["hist"]
    cur = m["neg"]["p"]
    by_code = {i["code"]: i for i in a["issues"]["issues"]}
    plan_codes = sorted(
        {c for t in backlog.tickets if t.track == "product" and t.horizon in ("now", "next") for c in t.issues}
        & set(by_code)
    )
    attributed = sum(by_code[c]["attributed_year"] for c in plan_codes)
    f = attributed / max(a["issues"]["year_neg_n"], 1)
    cost = sum(by_code[c]["rating_cost"] for c in plan_codes)
    bench = a["listing"]["benchmarks"][0] if a["listing"]["benchmarks"] else None
    return {
        "current": {"neg": cur, "rating": m["score"], "model": store_rating_if(cur, hist)},
        "plan_codes": plan_codes,
        "plan_share": f,
        "plan_cost": cost,
        "written_mean": a["windows"]["year"]["mean"]["mean"],
        "half": {"neg": cur * (1 - f / 2), "rating": store_rating_if(cur * (1 - f / 2), hist)},
        "full": {"neg": cur * (1 - f), "rating": store_rating_if(cur * (1 - f), hist)},
        "parity": {"name": bench["name"], "neg": bench["neg"]["p"], "rating": store_rating_if(bench["neg"]["p"], hist)}
        if bench
        else None,
    }


# ── context ───────────────────────────────────────────────────────────────────────────────────
def build_context(a: dict[str, Any], tax: Taxonomy, backlog: Backlog, df: pd.DataFrame | None = None) -> dict[str, Any]:
    """Everything the outputs need. `df` (the review frame) adds keyword-matched quotes and distinct
    reply-backlog counts; without it, quotes fall back to the analysis examples."""
    by_code = {i["code"]: i for i in a["issues"]["issues"]}
    watch = [v for v in a["versions"] if v["flag"] in ("worse", "watch")][-1:]
    used: set[str] = set()  # never show the same quote on two tickets
    tickets = []
    for t in backlog.tickets:
        codes = [c for c in t.issues if c in by_code]
        main = by_code[codes[0]] if codes else None
        base = (t.baseline or {}).get("share")
        n_now = main["recent"]["n"] if main else 0
        candidates: list[dict[str, Any]] = []
        if df is not None and (t.quote_terms or t.evidence == "release"):
            version = watch[0]["version"] if t.evidence == "release" and watch else None
            candidates = quotes_matching(df, t.quote_terms, codes=codes[:1] or None, version=version, k=8)
        candidates += main["examples"] if main else []
        q: list[dict[str, Any]] = []
        for x in candidates:
            if x["text"] not in used and len(q) < 2:
                q.append(x)
                used.add(x["text"])
        tickets.append(
            {
                **t.model_dump(),
                "labels": [tax.issues[c].label for c in codes],
                "evidence_lines": ticket_evidence(t, a, tax),
                "quotes": [_short(x) for x in q[:2]],
                "cost": sum(by_code[c]["rating_cost"] for c in codes),
                "main": main,
                "detectable": detectable_share(base, n_now) if base is not None and n_now else None,
            }
        )
    y, yb = a["windows"]["year"], a["windows"]["year_before"]
    yoy_p = two_prop_p(y["neg"]["k"], y["neg"]["n"], yb["neg"]["k"], yb["neg"]["n"])
    m = a["listing"]["main"]
    neg_n = a["issues"]["year_neg_n"]
    zero = {"k": 0, "n": neg_n, "p": 0.0, "lo": 0.0, "hi": 0.0}
    return {
        # Root-cause and target shares with every key present, so sparse data still renders.
        "targets": {k: a["issues"]["targets"].get(k, zero) for k in ("app", "service", "both", "none")},
        "roots": {k: a["issues"]["roots"].get(k, zero) for k in ("defect", "design", "policy", "service", "gap")},
        "tickets": tickets,
        "by_id": {t["id"]: t for t in tickets},
        "horizons": HORIZONS,
        "scen": scenarios(a, backlog),
        "replies_story": reply_story(a),
        "yoy_p": yoy_p,
        "one_star_written": a["gap"]["written"][0] / max(m["hist"][0], 1) if a["gap"] else None,
        "top": a["issues"]["issues"][0],
        "emerging": {e["code"]: e for e in a["emerging"]},
        "watch": watch,
        "playbook": playbook_rows(a, tax, backlog, df),
        "uncovered": plan_check(a, backlog),
        "plan_as_of": min((t.baseline["as_of"] for t in backlog.tickets if t.baseline), default=None),
    }


def plan_check(a: dict[str, Any], backlog: Backlog, top_n: int = 5) -> list[dict[str, str]]:
    """Signals in today's data that no ticket covers: an emerging or rising issue, or a top
    rating-cost issue. On a scheduled refresh this is how the page says the plan needs updating."""
    covered = {c for t in backlog.tickets for c in t.issues}
    days = a["windows"]["recent"]["days"]
    out: list[dict[str, str]] = []
    seen: set[str] = set()

    def add(code: str, label: str, why: str) -> None:
        if code not in covered and code not in seen:
            seen.add(code)
            out.append({"code": code, "label": label, "why": why})

    for e in a["emerging"]:
        add(
            e["code"],
            e["label"],
            f"emerging: {_p(e['last30']['p'])} of 1–2★ reviews in the last 30 days vs {_p(e['base']['p'])} before ({_pv(e['p'])})",
        )
    for i in a["issues"]["issues"]:
        if i["trend"] == "rising":
            add(
                i["code"],
                i["label"],
                f"rising: {_p(i['recent']['p'])} of 1–2★ reviews in the last {days} days vs {_p(i['prior']['p'])} ({_pv(i['trend_p'])})",
            )
    for i in a["issues"]["issues"][:top_n]:
        add(i["code"], i["label"], f"rating cost rank #{i['rank']} (−{i['rating_cost']:.3f}★)")
    return out


def _short(q: dict[str, Any], limit: int = 220) -> dict[str, Any]:
    text = re.sub(r"\s+", " ", q["text"]).strip()
    if len(text) > limit:
        text = text[:limit].rsplit(" ", 1)[0].rstrip(",.;:") + " …"
    return {**q, "text": text}


def playbook_rows(a: dict[str, Any], tax: Taxonomy, backlog: Backlog, df: pd.DataFrame | None) -> list[dict[str, Any]]:
    by_code = {i["code"]: i for i in a["issues"]["issues"]}
    days = a["windows"]["recent"]["days"]
    rows = []
    for r in backlog.replies:
        codes = [c for c in r.issues if c in by_code]
        rows.append(
            {
                **r.model_dump(),
                "labels": [tax.issues[c].label for c in r.issues],
                # Distinct reviews when the review frame is available; otherwise an upper bound.
                "unanswered": unanswered_negatives(df, codes, days)
                if df is not None
                else sum(by_code[c]["recent_unanswered"] for c in codes),
            }
        )
    return rows


def rec_charts(a: dict[str, Any], tax: Taxonomy) -> dict[str, str]:
    c = build_charts(a, tax)
    top = [i for i in a["issues"]["issues"] if i["rating_cost"] > 0][:12]
    c["cost"] = charts.hbar_chart(
        [
            Row(
                i["label"],
                i["rating_cost"],
                color_var=ROOT_VARS[i["root"]],
                tip=(
                    f"{i['label']}: −{i['rating_cost']:.3f}★ on the written-review average over "
                    f"{a['windows']['year']['days']} days · {i['recent']['p']:.0%} of recent 1–2★ reviews · owner {i['owner']}"
                ),
            )
            for i in top
        ],
        fmt=lambda v: f"−{v:.3f}★" if v else "0",
        title="Stars of written-review average lost to each issue",
    )
    return c


# ── outputs ───────────────────────────────────────────────────────────────────────────────────
def render_recommendation(a: dict[str, Any], tax: Taxonomy, ctx: dict[str, Any]) -> str:
    env = _env()
    return env.get_template("recommendation.html.j2").render(
        a=a, c=rec_charts(a, tax), tax=tax, root_vars=ROOT_VARS, **ctx
    )


def ticket_description_text(t: dict[str, Any], sep: str = "\n") -> str:
    parts = [f"Problem: {t['problem']}", "Evidence:", *[f"- {x}" for x in t["evidence_lines"]]]
    for q in t["quotes"]:
        parts.append(f'- Review ({q["stars"]}★, {q["date"]}): "{q["text"]}"')
    parts += ["Proposed change:", *[f"- {x}" for x in t["change"]]]
    parts += ["Acceptance criteria:", *[f"- {x}" for x in t["acceptance"]]]
    parts += [f"Owner: {t['owner']}", f"Effort (estimate): {t['effort']}", f"Success metric: {t['metric']}"]
    parts += [f"Confirm with (internal data): {t['confirm_with']}", f"Expected effect: {t['expected']}"]
    return sep.join(parts)


def jira_description(t: dict[str, Any], a: dict[str, Any]) -> str:
    def bullets(xs: list[str]) -> str:
        return "\n".join(f"* {x}" for x in xs)

    quotes = "\n".join(f"{{quote}}{q['text']} ({q['stars']}★, {q['date']}){{quote}}" for q in t["quotes"])
    base = t.get("baseline") or {}
    tracker = (
        f"\nFix tracker baseline: {base['share']:.1%} of 1–2★ reviews ({base['k']} of {base['n']}, "
        f"{base['days']} days to {base['as_of']})."
        if base
        else ""
    )
    return (
        f"h3. Problem\n{t['problem']}\n\n"
        f"h3. Evidence (Play Store reviews to {a['data_through']})\n{bullets(t['evidence_lines'])}\n{quotes}\n\n"
        f"h3. Proposed change\n{bullets(t['change'])}\n\n"
        f"h3. Acceptance criteria\n{bullets(t['acceptance'])}\n\n"
        f"*Owner:* {t['owner']}\n*Effort (estimate):* {t['effort']}\n"
        f"*Success metric:* {t['metric']}{tracker}\n*Confirm with (internal data):* {t['confirm_with']}\n"
        f"*Expected effect:* {t['expected']}"
        + (f"\n*Depends on:* {', '.join(t['depends_on'])}" if t["depends_on"] else "")
    )


def write_jira_csv(path: Path, a: dict[str, Any], ctx: dict[str, Any]) -> None:
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        # Jira's CSV importer maps repeated "Labels" columns to several labels.
        w.writerow(["Summary", "Issue Type", "Priority", "Labels", "Labels", "Labels", "Description"])
        for t in ctx["tickets"]:
            w.writerow(
                [
                    f"[{t['id']}] {t['title']}",
                    t["type"],
                    t["priority"],
                    "play-rating",
                    f"horizon-{t['horizon']}",
                    f"track-{t['track']}",
                    jira_description(t, a),
                ]
            )


def write_xlsx(path: Path, a: dict[str, Any], ctx: dict[str, Any]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Backlog"
    head = [
        ("ID", 9),
        ("When", 16),
        ("Track", 10),
        ("Priority", 10),
        ("Type", 8),
        ("Title", 42),
        ("Owner", 26),
        ("Effort (estimate)", 24),
        ("Issues", 28),
        ("Evidence", 70),
        ("Problem", 50),
        ("Proposed change", 70),
        ("Acceptance criteria", 55),
        ("Success metric", 40),
        ("Confirm with (internal data)", 36),
        ("Expected effect", 45),
        ("Baseline share", 10),
        ("Baseline n", 10),
        ("Depends on", 12),
    ]
    ws.append([h for h, _ in head])
    for i, (_, wdt) in enumerate(head, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    for t in ctx["tickets"]:
        base = t.get("baseline") or {}
        ws.append(
            [
                t["id"],
                HORIZONS[t["horizon"]][0],
                "Product" if t["track"] == "product" else "Operations",
                t["priority"],
                t["type"],
                t["title"],
                t["owner"],
                t["effort"],
                "\n".join(t["labels"]),
                "\n".join(
                    [f"• {x}" for x in t["evidence_lines"]]
                    + [f"“{q['text']}” ({q['stars']}★, {q['date']})" for q in t["quotes"]]
                ),
                t["problem"],
                "\n".join(f"• {x}" for x in t["change"]),
                "\n".join(f"• {x}" for x in t["acceptance"]),
                t["metric"],
                t["confirm_with"],
                t["expected"],
                base.get("share"),
                base.get("n"),
                ", ".join(t["depends_on"]),
            ]
        )
    _style_sheet(ws, Font, PatternFill, Alignment, pct_cols=[17])

    ev = wb.create_sheet("Issue evidence")
    days = a["windows"]["recent"]["days"]
    ycols = [
        ("Rank", 6),
        ("Issue", 36),
        ("Owner", 28),
        ("Root cause", 26),
        ("Rating cost (★)", 12),
        (f"Share of 1–2★, last {days}d", 12),
        ("95% CI low", 10),
        ("95% CI high", 10),
        ("1–2★ reviews (n)", 10),
        (f"Share, previous {days}d", 12),
        ("Trend", 9),
        ("Trend p", 8),
        ("Share, 12 months", 10),
        ("Emerging (last 30d)", 12),
        ("Unanswered 1–2★, last 90d", 12),
    ]
    ev.append([h for h, _ in ycols])
    for i, (_, wdt) in enumerate(ycols, 1):
        ev.column_dimensions[get_column_letter(i)].width = wdt
    em = ctx["emerging"]
    for i in a["issues"]["issues"]:
        ev.append(
            [
                i["rank"],
                i["label"],
                i["owner"],
                i["root_label"],
                round(i["rating_cost"], 4),
                i["recent"]["p"],
                i["recent"]["lo"],
                i["recent"]["hi"],
                i["recent"]["n"],
                i["prior"]["p"],
                i["trend"],
                round(i["trend_p"], 3),
                i["year"]["p"],
                em[i["code"]]["last30"]["p"] if i["code"] in em else None,
                i["recent_unanswered"],
            ]
        )
    _style_sheet(ev, Font, PatternFill, Alignment, pct_cols=[6, 7, 8, 10, 13, 14])

    about = wb.create_sheet("About")
    for row in [
        ["MoveInSync employee app — Play Store rating backlog"],
        [f"Evidence: Google Play reviews up to {a['data_through']}; generated {a['generated_at']}."],
        ["Share = share of 1–2★ written reviews in the window that name the issue (a review can name up to 3)."],
        [
            "Rating cost = stars of the 12-month written-review average lost to the issue (upper bound of what a fix recovers)."
        ],
        [
            "Baseline = share when the recommendation was made; the monitor's Fix tracker compares the latest share with it."
        ],
        ["Effort is an estimate to be confirmed by the owning team."],
        ["Method and accuracy: docs/METHOD.md. Regenerate: uv run ratings recommend."],
    ]:
        about.append(row)
    about.column_dimensions["A"].width = 120
    about["A1"].font = Font(bold=True, size=13)
    wb.save(path)


def _style_sheet(ws: Any, Font: Any, PatternFill: Any, Alignment: Any, pct_cols: list[int]) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F0EFEC")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column in pct_cols and isinstance(cell.value, float):
                cell.number_format = "0.0%"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def render_playbook(a: dict[str, Any], tax: Taxonomy, ctx: dict[str, Any]) -> str:
    env = _env()
    env.autoescape = False  # Markdown, not HTML
    return env.get_template("reply_playbook.md.j2").render(a=a, tax=tax, **ctx)


def write_all(
    out_dir: Path, a: dict[str, Any], tax: Taxonomy, backlog: Backlog, df: pd.DataFrame | None = None
) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    ctx = build_context(a, tax, backlog, df)
    paths = {
        "recommendation": out_dir / "recommendation.html",
        "xlsx": out_dir / "backlog.xlsx",
        "jira": out_dir / "backlog_jira.csv",
        "playbook": out_dir / "reply_playbook.md",
    }
    paths["recommendation"].write_text(render_recommendation(a, tax, ctx), encoding="utf-8")
    write_xlsx(paths["xlsx"], a, ctx)
    write_jira_csv(paths["jira"], a, ctx)
    paths["playbook"].write_text(render_playbook(a, tax, ctx), encoding="utf-8")
    return paths
